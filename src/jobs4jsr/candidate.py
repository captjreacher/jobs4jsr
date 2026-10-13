import hashlib, json, re
from pathlib import Path
from pypdf import PdfReader
from docx import Document

def extract(path):
    path = Path(path)
    if path.suffix.lower() == '.pdf':
        return '\n'.join(p.extract_text() or '' for p in PdfReader(path).pages)
    if path.suffix.lower() == '.docx':
        d = Document(path)
        return '\n'.join([p.text for p in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
    return path.read_text(encoding='utf-8-sig')

def build(source, destination):
    destination = Path(destination); destination.mkdir(parents=True, exist_ok=True)
    docs = {p.name: extract(p) for p in Path(source).iterdir() if p.suffix.lower() in {'.pdf','.docx','.txt'}}
    if len(docs) < 4: raise ValueError('At least two CVs and two templates are required')
    (destination/'extracted_documents.json').write_text(json.dumps(docs,indent=2),encoding='utf-8')
    reviewed = destination/'reviewed_profile.json'
    if not reviewed.exists():
        draft = dict(name=None, email=None, skills=[], experience=[], qualifications=[], availability=None, work_rights=None, barista=False, table_service=False, uncertainties=['Human source review required'],style=dict(greeting='Dear Recruitment Team',closing='Yours sincerely',tone='Warm, direct and concise'))
        (destination/'profile_draft.json').write_text(json.dumps(draft,indent=2),encoding='utf-8')
        raise ValueError('Documents extracted locally. Review extracted_documents.json and fill reviewed_profile.json using profile_draft.json; then rerun profile. No experience will be inferred automatically.')
    profile=json.loads(reviewed.read_text(encoding='utf-8'))
    if not profile.get('name') or not profile.get('experience') or not profile.get('skills'): raise ValueError('Reviewed profile is incomplete')
    profile['sources']=[dict(file=k,sha256=hashlib.sha256(Path(source,k).read_bytes()).hexdigest()) for k in docs]
    (destination/'jess_profile.json').write_text(json.dumps(profile,indent=2),encoding='utf-8')
    summary='# Verified candidate summary\n\n'+ '\n'.join(f"- {e['role']}, {e['employer']}, {e.get('start','Unknown')} to {e.get('end','Unknown')}: {', '.join(e['duties'])}." for e in profile['experience'])
    summary+='\n\nSkills: '+', '.join(profile['skills'])+'.\nQualifications: '+ '; '.join(profile['qualifications'])+'.\n\nUncertainties:\n'+'\n'.join('- '+x for x in profile['uncertainties'])
    (destination/'candidate_summary.md').write_text(summary,encoding='utf-8')
    return profile

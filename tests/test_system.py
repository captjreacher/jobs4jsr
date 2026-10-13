import json,tempfile,unittest
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from jobs4jsr import candidate,database,deduplication,letters,reporting,scoring,scheduler

class SystemTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.root=Path(self.temp.name)
        self.p={'name':'Test Candidate','skills':['customer service'],'experience':[{'role':'Retail Assistant','employer':'Example Store','duties':['welcomed customers','processed sales']}],'style':{'greeting':'Dear Recruitment Team','closing':'Yours sincerely'}}
        self.j=dict(vacancy_id='123',source='Test direct employer',id=1,employer='Test Venue',title='Front of House Team Member',location='Viaduct Harbour',distance=1.1,description='Customer service and teamwork. Training provided; experience not necessary.',url='https://example.org/jobs/123?utm_source=test',posted='2026-10-12',recent=True,job_type='PART_TIME',found='2026-10-14',match=100,why='Customer service',status='New')
    def tearDown(self): self.temp.cleanup()
    def test_distance(self):
        self.assertAlmostEqual(scoring.distance((-36.8485,174.7633),(-36.8485,174.7633)),0)
        self.assertAlmostEqual(scoring.distance((0,0),(0,1)),111.195,places=2)
        self.assertGreater(scoring.distance((-36.8485,174.7633),(-37.0,174.8)),10)
    def test_score(self):
        self.assertEqual(scoring.score(self.j,self.p)[0],90)
        self.assertIsNone(scoring.score(dict(self.j,distance=11),self.p))
        self.assertIsNone(scoring.score(dict(self.j,title='Senior Waiter'),self.p))
        self.assertIsNone(scoring.score(dict(self.j,description='Previous experience in similar role within food & beverage'),self.p))
    def test_dedup(self):
        self.assertEqual(deduplication.canonical('https://example.org/jobs/123?utm_source=test#top'),'https://example.org/jobs/123')
        self.assertTrue(deduplication.duplicate(self.j,dict(self.j,url='https://other.org/job/1')))
        self.assertFalse(deduplication.duplicate(dict(self.j,url='https://other.org/job/2',posted='2026-12-13',vacancy_id='124'),self.j))
        self.assertTrue(deduplication.duplicate(self.j,dict(self.j,url='https://board.org/456',source='Different board',posted='2026-08-01',vacancy_id='456')))
    def test_extraction(self):
        p=self.root/'a.txt'; p.write_text('Verified source'); self.assertEqual(candidate.extract(p),'Verified source')
        d=Document();d.add_paragraph('DOCX source');d.save(self.root/'a.docx');self.assertEqual(candidate.extract(self.root/'a.docx'),'DOCX source')
    def test_pdf_extraction(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        w=PdfWriter(); page=w.add_blank_page(width=400,height=400)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 10 350 Td (Verified PDF evidence) Tj ET')
        page[NameObject('/Contents')]=stream
        pdf=self.root/'sample.pdf'
        with pdf.open('wb') as f: w.write(f)
        self.assertIn('Verified PDF evidence',candidate.extract(pdf))
    def test_letters_and_reports(self):
        md,word=letters.generate(self.j,self.p,self.root,'2026-10-14'); self.j.update(letter_md=md,letter_docx=word)
        self.assertIn('Test Venue', (self.root/md).read_text())
        self.assertNotIn('work visa',(self.root/md).read_text())
        self.assertGreater(len(Document(self.root/word).paragraphs),6)
        latest=reporting.report(self.root,'2026-10-14',[self.j],['Source failed <example>'],'test')
        import re
        for file in (latest,self.root/'outputs/reports/2026-10-14/report.html'):
            for link in re.findall('href="([^"]+)"',file.read_text()):
                if not link.startswith('http'): self.assertTrue((file.parent/link).exists(),link)
        self.assertIn('&lt;example&gt;',latest.read_text())
    def test_database_status(self):
        c=database.connect(self.root/'history.sqlite'); ident=database.add(c,self.j); c.commit()
        database.status(c,ident,'Applied'); self.assertEqual(database.jobs(c)[0]['status'],'Applied')
        self.assertEqual(c.execute('select count(*) from status_history').fetchone()[0],1)
        with self.assertRaises(ValueError): database.status(c,99,'Applied')
        c.close()
    def test_schedule(self):
        zone=ZoneInfo('Pacific/Auckland')
        self.assertEqual(scheduler.due_slot(datetime(2026,10,14,8,tzinfo=zone)),'2026-10-12-scheduled')
        self.assertEqual(scheduler.due_slot(datetime(2026,10,14,9,tzinfo=zone)),'2026-10-14-scheduled')
        self.assertEqual(scheduler.due_slot(datetime(2026,10,18,9,tzinfo=zone)),'2026-10-16-scheduled')
    def test_cli_scheduled_idempotency(self):
        import sys
        from unittest.mock import patch
        from jobs4jsr import cli
        (self.root/'config').mkdir(); (self.root/'config/settings.example.json').write_text('{}')
        (self.root/'data/private').mkdir(parents=True)
        (self.root/'data/private/jess_profile.json').write_text(json.dumps(self.p))
        with patch.object(sys,'argv',['jobs4jsr','--root',str(self.root),'scheduled']), patch('jobs4jsr.cli.discovery.discover',return_value=([self.j],[])) as discover:
            cli.main(); cli.main(); self.assertEqual(discover.call_count,1)
        c=database.connect(self.root/'data/private/jobs.sqlite3')
        self.assertEqual(len(database.jobs(c)),1); c.close()
    def test_source_failure_retry_and_empty_report(self):
        import sys
        from unittest.mock import patch
        from jobs4jsr import cli
        (self.root/'config').mkdir(); (self.root/'config/settings.example.json').write_text('{}')
        (self.root/'data/private').mkdir(parents=True)
        (self.root/'data/private/jess_profile.json').write_text(json.dumps(self.p))
        with patch.object(sys,'argv',['jobs4jsr','--root',str(self.root),'scheduled']), patch('jobs4jsr.cli.discovery.discover',return_value=([],['SOURCE FAILED example: timeout'])) as discover:
            for _ in range(2):
                with self.assertRaises(SystemExit) as caught: cli.main()
                self.assertEqual(caught.exception.code,2)
            self.assertEqual(discover.call_count,2)
        self.assertIn('SOURCE FAILED',(self.root/'outputs/reports/latest.html').read_text())
    def test_profile_review_gate(self):
        source=self.root/'source';source.mkdir()
        for n in range(4): (source/f'doc{n}.txt').write_text('Evidence without an approved profile')
        with self.assertRaises(ValueError): candidate.build(source,self.root/'private')
        self.assertTrue((self.root/'private/profile_draft.json').exists())
        self.assertFalse((self.root/'private/jess_profile.json').exists())
    def test_run_uniqueness(self):
        import sqlite3
        c=database.connect(self.root/'db.sqlite'); c.execute("insert into runs(id) values('slot')");c.commit()
        with self.assertRaises(sqlite3.IntegrityError): c.execute("insert into runs(id) values('slot')")
        c.close()

if __name__=='__main__': unittest.main()
